# -*- coding: utf-8 -*-
"""WUZHEN — Waterways & Bridges  ·  6-slide condensed deck.

Compressed from the 13-slide version. One screen = one idea, Apple-style:
black / white / light-gray rhythm, oversized titles, photos dominant, short
copy, generous margins. Text boxes have auto-grow OFF so every box keeps the
geometry written here — overflow is caught by tools/audit_layout.py instead of
silently pushing into a neighbour.

All numbers, names, image paths and credits come from the project site
(chapters.js, Facts, WhyWater, Timeline, Atlas, Gallery, team.js,
images-sources.json). Nothing is invented.

Run:  cd wuzhen-project && python tools/build_wuzhen_pptx_6p.py
Out:  wuzhen_presentation_6p.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from wuzhen_pptx_helpers import (
    INK, SUB, BG_L, WHITE, WGREEN, DEEPWATER, HAIR, SOFT,
    add_text, add_rect, add_pic, set_bg, transition, animations, add_notes,
)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG = os.path.join(ROOT, "public", "images")
CROP = os.path.join(ROOT, "tools", "_pptx_crops")
OUT = os.path.join(ROOT, "wuzhen_presentation_6p.pptx")
os.makedirs(CROP, exist_ok=True)

PX = 144.0                                  # 1920px / 13.333in
I = lambda px: px / PX                      # design px -> inches
T = lambda size, color, bold=False, spc=None: \
    dict(size=size, color=color, bold=bold, spc=spc)
SW, SH = 13.3333, 7.5


def crop(src, w, h, out, pos="center"):
    """Cover-crop a source photo into an exact pixel box (no stretching)."""
    from PIL import Image
    dst = os.path.join(CROP, out)
    if os.path.exists(dst) and os.path.getsize(dst) > 0:
        return dst
    with Image.open(os.path.join(IMG, src)) as im:
        imw, imh = im.size
        tr, sr = w / h, imw / imh
        if sr > tr:
            nw = round(imh * tr)
            frac = 0.5 if pos == "center" else pos
            x0 = round((imw - nw) * frac)
            im = im.crop((x0, 0, x0 + nw, imh))
        elif sr < tr:
            nh = round(imw / tr)
            frac = 0.5 if pos == "center" else (0.30 if pos == "top" else pos)
            y0 = round((imh - nh) * frac)
            im = im.crop((0, y0, imw, y0 + nh))
        if im.size != (w, h):
            im = im.resize((w, h), Image.LANCZOS)
        im.convert("RGB").save(dst, "JPEG", quality=92)
    return dst


def veil(slide, y, h, alpha=0.55, color=None, x=0, w=1920):
    """Dark scrim over a photo — scrims keep type legible, no blur needed."""
    return add_rect(slide, I(x), I(y), I(w), I(h),
                    fill=color or INK, fill_alpha=alpha)


def pic(slide, src, x, y, w, h, out=None, pos="center"):
    return add_pic(slide, crop(src, w, h, out or f"{src}__{w}x{h}"),
                   I(x), I(y), I(w), I(h), crop="cover")


def line(slide, x, y, w, color=DEEPWATER, h=3):
    return add_rect(slide, I(x), I(y), I(w), I(h), fill=color)


def vrule(slide, x, y, h, color=DEEPWATER, w=2):
    """Vertical hairline: x/y position, h = height, w = thickness."""
    return add_rect(slide, I(x), I(y), I(w), I(h), fill=color)


TEAM = [
    ("蒋盛熠", "JIANG SHENGYI", "Project lead",
     "Coordinated the four-week plan, kept the group on schedule, and led the final presentation.",
     "team-m2-av-200x200-jpg.jpg"),
    ("汪瀚宇", "WANG HANYU", "Research & mapping",
     "Traced the canal routes and surveyed the bridges, mapping how the waterways connect.",
     "team-m1-av-200x200-jpg.jpg"),
    ("鲁昂", "LU ANG", "Design & web",
     "Built the interactive site — layout, motion and the visual story — and managed the image assets.",
     "team-m3-av-200x200-jpg.jpg"),
    ("朱钟乐", "ZHU ZHONGLE", "Photography",
     "Shot the on-site photo tour of the canals and bridges, and sourced the images.",
     "team-m4-av-200x200-jpg.jpg"),
    ("沈毅程", "SHEN YICHENG", "Copy & research",
     "Wrote the on-screen text, checked the bridge facts and history, and kept the story accurate.",
     "team-m5-av-200x200-jpg.jpg"),
    ("沈煜程", "SHEN YUCHENG", "Production & final review",
     "Ran the weekly checklists, prepared the slides and materials, and did the final self-review.",
     "team-m6-av-200x200-jpg.jpg"),
]

NOTES = [
"OPENING — 'Wuzhen, a water town shaped by water.' Then ask: 'Our question is: "
"why is water the main line of the ancient town?' Everything in this deck answers "
"that one question. Keep it calm and slow — this is an exhibition, not a report.",
"Visual research. 'We looked at Wuzhen as a water system, not just a picture. "
"These are the three things we kept returning to.' Let the three photos speak for "
"a moment, then read the website line once. Simple and confident.",
"The waterway atlas. 'This is not a map. It is our attempt to trace the route: "
"north gate to the early market.' Four nodes, no labels we invented. The line is "
"what connects everything we saw.",
"Answer one and answer two. Why? Because water decides where houses are built, "
"where bridges must be built, and how streets have to be laid. That is the "
"argument. Then Facts — and only the numbers we actually have: 72 ancient stone "
"bridges, and nearly 10,000 metres of waterway in Xizha, in a cross-shaped "
"system. Read them slowly. We did not guess extra numbers.",
"Past to present. Day to night. 'In the past the waterway was transport and "
"trade. Today it is tourism and lanterns.' Show the day photo, then the night "
"photo, and only then move to the words on the right.",
"Closing. 'Wuzhen, water connects everything.' Thank the team. 'Water carries "
"the town, carries our research, and carries this project.' One breath, then stop.",
]

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def new(bg):
    s = prs.slides.add_slide(BLANK)
    set_bg(s, bg)
    return s


# ============================================================================
# 1 — COVER + OUR QUESTION
# ============================================================================
s = new(INK)
pic(s, "hero_wuzhen-1920-jpg.jpg", 0, 0, 1920, 1080, "cv_hero")
veil(s, y=0, h=1080, alpha=0.40)
veil(s, y=0, h=560, alpha=0.60)
veil(s, y=620, h=460, alpha=0.70)

add_text(s, I(72), I(58), I(1100), I(30),
         [("GROUP 90706  ·  6 MEMBERS  ·  PBL PROJECT  ·  2026",
           T(12.5, SOFT, True, 170))])
title_s = add_text(s, I(72), I(272), I(1100), I(150),
                   [("WUZHEN", T(128, WHITE, True, -200))])
sub_s = add_text(s, I(78), I(506), I(900), I(36),
                 [("WATERWAYS  &  BRIDGES", T(26, WHITE, True, 230))])
tag_s = add_text(s, I(78), I(568), I(900), I(32),
                 [("A water town shaped by water.", T(21, SOFT))])
line(s, 1140, 700, 90, WGREEN, 3)
add_text(s, I(1140), I(734), I(708), I(26),
         [("OUR QUESTION", T(13, WHITE, True, 230))])
q = add_text(s, I(1140), I(776), I(708), I(170),
             [("Why is water the main line", T(32, WHITE, True, -40)),
              ("of the ancient town?", T(32, WHITE, True, -40))],
             line_spacing=1.14)
add_text(s, I(1140), I(946), I(708), I(30),
         [("为什么水是古镇的主线？", T(19, SOFT))])
add_text(s, I(72), I(1000), I(900), I(28),
         [("Everything in this deck answers one question.", T(14, SOFT, False, 60))])
transition(s, "black", 1400)
animations(s, [(title_s, "fade", 100, 800),
               (sub_s, "fade", 300, 700),
               (q, "fade", 900, 800)])
add_notes(s, NOTES[0])

# ============================================================================
# 2 — VISUAL RESEARCH  (white)
# ============================================================================
s = new(WHITE)
add_text(s, I(72), I(56), I(1100), I(28),
         [("02 — WHAT WE SAW", T(12.5, WGREEN, True, 230))])
vis_title = add_text(s, I(72), I(94), I(1776), I(84),
                     [("VISUAL RESEARCH", T(60, INK, True, -110))])
add_text(s, I(72), I(194), I(1100), I(30),
         [("Wuzhen as a water system, not just a picture.", T(20, SUB))])
line(s, 72, 246, 1776, HAIR, 1)

SHOTS = [
    ("waterway-jpg.jpg", "WATER — THE CROSS-SHAPED SYSTEM", "水：把古镇分成不同区域"),
    ("bridge2-jpg.jpg", "BRIDGES — A SHORT ARC LINKS TWO BANKS", "桥：一段短拱连接两岸"),
    ("night-jpg.jpg", "NIGHT — LANTERNS POUR LIGHT ONTO WATER", "夜晚：灯光落进水里"),
]
cards = []
for i, (src, en, zh) in enumerate(SHOTS):
    x = 72 + i * 608
    cards.append(pic(s, src, x, 268, 560, 336))
    add_text(s, I(x), I(628), I(540), I(26),
             [(en, T(13.5, INK, True, 60))])
    add_text(s, I(x), I(662), I(540), I(28),
             [(zh, T(15, SUB))])

# website presentation band
add_rect(s, I(72), I(744), I(1776), I(258), fill=BG_L)
add_rect(s, I(72), I(744), I(1776), I(40), fill=HAIR)
for k in range(3):
    add_rect(s, I(88 + k * 18), I(758), I(10), I(10), fill=SOFT)
add_rect(s, I(300), I(752), I(1600), I(18), fill=WHITE)
add_text(s, I(312), I(753), I(900), I(18),
         [("wuzhen-project.github.io/wuzhen-waterways/", T(10.5, SUB, True))])
add_text(s, I(300), I(808), I(1560), I(36),
         [("One website, seven chapters, every image credited.",
           T(24, INK, True, -20))])
add_text(s, I(300), I(862), I(1560), I(80),
         [("01 Facts · 02 Gallery · 03 Keywords · 04 Why Water · "
           "05 Timeline · 06 Sources · 07 Team", T(15, SUB)),
          ("The exhibition is digital, public, and browsable.", T(15, SUB)),
          ("Every photo carries its source and licence.", T(15, SUB))],
         line_spacing=1.28)
add_text(s, I(300), I(968), I(1560), I(24),
         [("Built for this PBL project — live and open.", T(13, SOFT, False, 60))])
transition(s, "fade", 1000)
animations(s, [(cards[0], "zoom", 0, 700),
               (cards[1], "zoom", 0, 700),
               (cards[2], "zoom", 0, 700),
               (s.shapes[-5], "fade", 400, 600)])
add_notes(s, NOTES[1])

# ============================================================================
# 3 — WATERWAY ATLAS  (black)
# ============================================================================
s = new(INK)
add_text(s, I(72), I(56), I(1100), I(28),
         [("03 — WATERWAY ATLAS", T(12.5, WGREEN, True, 230))])
atlas_title = add_text(s, I(72), I(94), I(1100), I(84),
                       [("WATERWAY ATLAS", T(60, WHITE, True, -110))])
atlas_sub = add_text(s, I(72), I(194), I(1100), I(30),
                     [("Not a map. An attempt to trace the line.", T(20, SOFT))])
add_text(s, I(1340), I(194), I(508), I(30),
         [("NOT A MAP", T(12.5, WGREEN, True, 200))], align=PP_ALIGN.RIGHT)

PTS = [(250, 520), (750, 600), (1250, 520), (1750, 600)]
for a, b in zip(PTS, PTS[1:]):
    x1, y1 = a
    x2, y2 = b
    line(s, x1, min(y1, y2), x2 - x1, WGREEN, 2)
    vrule(s, x2 - 2, min(y1, y2), abs(y2 - y1), WGREEN, 3)
NODES = [
    (250, 520, "01", "NORTH WATER GATE", "北栅水门"),
    (750, 600, "02", "XIZHA BRIDGE", "西栅古桥"),
    (1250, 520, "03", "LANTERN WHARF", "灯影码头"),
    (1750, 600, "04", "MORNING MARKET", "清晨市集"),
]
for n, (x, y, num, en, zh) in enumerate(NODES):
    add_rect(s, I(x - 6), I(y - 6), I(12), I(12), fill=WGREEN)
    bx = x - 95
    by = y + 34
    add_text(s, I(bx), I(by), I(190), I(22),
             [(num, T(12, WGREEN, True, 200))])
    add_text(s, I(bx), I(by + 26), I(190), I(26),
             [(en, T(12, WHITE, True, 40))])
    add_text(s, I(bx), I(by + 56), I(190), I(28),
             [(zh, T(15, SOFT))])
atlas_rule = line(s, 72, 852, 1776, DEEPWATER, 2)
atlas_cap = add_text(s, I(72), I(888), I(1776), I(36),
                     [("Water is the route. Houses, bridges and streets follow it.",
                       T(24, WHITE, True, -20))])
add_text(s, I(72), I(946), I(1776), I(28),
         [("This line is our research attempt, not an official map.",
           T(14, SOFT, False, 60))])
transition(s, "fade", 1000)
animations(s, [(atlas_title, "fade", 0, 700),
               (atlas_sub, "fade", 0, 500),
               (atlas_rule, "fade", 0, 700),
               (atlas_cap, "fade", 0, 600)])
add_notes(s, NOTES[2])

# ============================================================================
# 4 — WHY WATER?  +  FACTS  (light gray)
# ============================================================================
s = new(BG_L)
add_text(s, I(72), I(56), I(1100), I(28),
         [("04 — WHY WATER?   /   THE ANSWER", T(12.5, WGREEN, True, 230))])
whyt = add_text(s, I(72), I(94), I(1100), I(84),
                [("WHY WATER?", T(60, INK, True, -110))])

bar_water = add_rect(s, I(72), I(212), I(360), I(62), fill=INK)
add_text(s, I(72), I(226), I(360), I(36),
         [("WATER", T(26, WHITE, True, 180))], align=PP_ALIGN.CENTER)
add_rect(s, I(452), I(212), I(280), I(62), fill=HAIR)
add_text(s, I(452), I(230), I(280), I(28),
         [("DECIDES", T(15, INK, True, 130))], align=PP_ALIGN.CENTER)
bar_out = add_rect(s, I(752), I(212), I(1096), I(62), fill=INK)
add_text(s, I(752), I(230), I(1096), I(28),
         [("HOUSE PLACES  ·  BRIDGE PLACES  ·  STREET LINES  ·  TRADITIONAL LIFE",
           T(15, WHITE, True, 70))], align=PP_ALIGN.CENTER)

why_copy = add_text(s, I(72), I(318), I(1776), I(36),
                    [("Because water comes first, everything else has to follow it.",
                      T(22, INK, True, -10))])
add_text(s, I(72), I(366), I(1776), I(28),
         [("水先存在，所以房屋、桥梁和街道都必须沿着它安排。", T(16, SUB))])

line(s, 72, 448, 1776, HAIR, 1)
add_text(s, I(72), I(478), I(1100), I(28),
         [("FACTS", T(13, WGREEN, True, 230))])
fact72 = add_text(s, I(72), I(540), I(830), I(120),
                  [("72", T(120, INK, True, -200))])
add_text(s, I(72), I(694), I(830), I(28),
         [("ANCIENT STONE BRIDGES", T(13.5, SUB, True, 100))])
add_text(s, I(72), I(728), I(830), I(30),
         [("现存古石桥，位于西栅", T(15, SUB))])
vrule(s, 986, 540, 218, HAIR, 2)
fact10k = add_text(s, I(1030), I(540), I(818), I(120),
                   [("10,000", T(80, INK, True, -200))])
add_text(s, I(1030), I(694), I(818), I(28),
         [("METRES OF WATERWAYS", T(13.5, SUB, True, 100))])
add_text(s, I(1030), I(728), I(818), I(30),
         [("近万米河道，内河水系呈十字形", T(15, SUB))])
line(s, 72, 800, 1776, HAIR, 1)
add_text(s, I(72), I(840), I(1776), I(36),
         [("Water divides the town into zones — and connects it too.",
           T(24, INK, True, -10))])
add_text(s, I(72), I(898), I(1776), I(28),
         [("Numbers from our field research and the project website.",
           T(14, SOFT, False, 60))])
transition(s, "fade", 1000)
animations(s, [(whyt, "fade", 0, 600),
               (bar_water, "fade", 0, 500),
               (bar_out, "fade", 0, 500),
               (why_copy, "fade", 0, 500),
               (fact72, "zoom", 0, 700),
               (fact10k, "zoom", 0, 700)])
add_notes(s, NOTES[3])

# ============================================================================
# 5 — LIFE / NIGHT   +   PAST -> PRESENT  (black)
# ============================================================================
s = new(INK)
day_pic = pic(s, "street-jpg.jpg", 0, 0, 1150, 540, "day_crop", "top")
night_pic = pic(s, "night-jpg.jpg", 0, 540, 1150, 540, "night_crop")
veil(s, y=0, h=420, alpha=0.58, x=0, w=1150)
veil(s, y=540, h=420, alpha=0.62, x=0, w=1150)
line(s, 0, 539, 1150, INK, 4)

day_label = add_text(s, I(56), I(146), I(1040), I(36),
                     [("BY DAY — TRAVEL · TRANSPORTATION · TRADE",
                       T(15, WHITE, True, 100))])
night_label = add_text(s, I(56), I(592), I(1040), I(36),
                       [("BY NIGHT — LANTERNS POUR LIGHT ONTO THE WATER",
                         T(15, WHITE, True, 100))])

add_text(s, I(1210), I(72), I(660), I(28),
         [("05 — HISTORY", T(12.5, WGREEN, True, 230))])
hist_title = add_text(s, I(1210), I(112), I(660), I(140),
                      [("PAST  →", T(44, WHITE, True, -70)),
                       ("PRESENT", T(44, WHITE, True, -70))], line_spacing=1.1)
line(s, 1210, 300, 660, DEEPWATER, 2)
add_text(s, I(1210), I(344), I(660), I(26),
         [("PAST", T(14, WGREEN, True, 200))])
hist_past = add_text(s, I(1210), I(382), I(660), I(110),
                     [("The waterway was travel, transport and trade.",
                       T(20, WHITE, True, -10)),
                      ("船是路，水是交通。", T(15, SOFT))], line_spacing=1.2)
hist_kick = add_text(s, I(1210), I(516), I(660), I(30),
                     [("→  WATER", T(15, WGREEN, True, 200))])
add_text(s, I(1210), I(570), I(660), I(26),
         [("PRESENT", T(14, WGREEN, True, 200))])
hist_present = add_text(s, I(1210), I(608), I(660), I(110),
                        [("Today it is tourism and lanterns.", T(20, WHITE, True, -10)),
                         ("今天，它变成旅游和夜晚的灯。", T(15, SOFT))], line_spacing=1.2)
add_text(s, I(1210), I(794), I(660), I(130),
         [("Water has always been the main line — only the role changed.",
           T(17, SOFT, False))], line_spacing=1.3)
transition(s, "fade", 1000)
animations(s, [(day_pic, "fade", 0, 600),
               (night_pic, "fade", 0, 600),
               (day_label, "fade", 0, 500),
               (night_label, "fade", 0, 500),
               (hist_title, "fade", 300, 700),
               (hist_past, "fade", 0, 600),
               (hist_kick, "fade", 0, 500),
               (hist_present, "fade", 0, 600)])
add_notes(s, NOTES[4])

# ============================================================================
# 6 — TEAM  +  CONCLUSION  (black)
# ============================================================================
s = new(INK)
add_text(s, I(72), I(56), I(1100), I(28),
         [("06 — OUR TEAM   /   CONCLUSION", T(12.5, WGREEN, True, 230))])
add_text(s, I(1200), I(56), I(648), I(28),
         [("Team members: six", T(13, SOFT, False, 100))], align=PP_ALIGN.RIGHT)

team_pic = pic(s, "team-group-4096x2048-jpg.jpg", 72, 106, 1776, 258, "team_crop")
veil(s, y=106, h=258, alpha=0.52, x=72, w=1776)

for i, (zh, en, role, contrib, photo) in enumerate(TEAM):
    x = 72 + i * 297
    pic(s, photo, x, 406, 100, 100)
    add_text(s, I(x), I(526), I(272), I(30),
             [(zh, T(21, WHITE, True))])
    add_text(s, I(x), I(562), I(272), I(24),
             [(en, T(12, SOFT, True, 90))])
    add_text(s, I(x), I(592), I(272), I(28),
             [(role, T(14, WHITE, True))], line_spacing=1.1)
    add_text(s, I(x), I(622), I(272), I(112),
             [(contrib, T(10, SOFT))], line_spacing=1.22)

line(s, 0, 742, 1920, DEEPWATER, 2)
conc_title = add_text(s, I(72), I(788), I(1100), I(100),
                      [("WUZHEN", T(88, WHITE, True, -160))])
conc_sub = add_text(s, I(78), I(908), I(1100), I(36),
                    [("Water connects everything.", T(26, WHITE, True, -10))])
add_text(s, I(78), I(958), I(1100), I(60),
         [("Water carries the town, carries our research, and carries this project.",
           T(14, SOFT, False))], line_spacing=1.3)
add_text(s, I(1200), I(800), I(648), I(26),
         [("WATERWAYS & BRIDGES  ·  GROUP 90706  ·  PBL 2026",
           T(11.5, SUB, True, 150))], align=PP_ALIGN.RIGHT)
add_text(s, I(1200), I(842), I(648), I(96),
         [("Image sources — Wikimedia Commons (韩笃一, Immanuel Giel, Jakub",
           T(10.5, SUB)),
          ("Hałun, Gerbil, Nico le terrible, Sastognuti, N509FZ,",
           T(10.5, SUB)),
          ("ChinaUli2010; CC0 / CC BY / CC BY-SA)", T(10.5, SUB))],
         align=PP_ALIGN.RIGHT, line_spacing=1.34)
add_text(s, I(1200), I(952), I(648), I(26),
         [("Full credits live on the project website, in Sources.",
           T(11, SOFT, False, 60))], align=PP_ALIGN.RIGHT)
transition(s, "black", 1500)
animations(s, [(team_pic, "zoom", 0, 700),
               (conc_title, "fade", 500, 700),
               (conc_sub, "fade", 0, 700)])
add_notes(s, NOTES[5])

# -----------------------------------------------------------------------------
prs.save(OUT)
print("OK ", OUT)
print("slides:", len(prs.slides._sldIdLst), " size:", SW, "x", SH, " (1920x1080)")
